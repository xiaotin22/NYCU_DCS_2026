# -*- coding: utf-8 -*-
"""Build an editable single-slide PPTX of the CA module signal-flow diagram.
Coordinates are taken 1:1 from the verified dark SVG (1280x720 px == 13.333x7.5in @ 96px/in)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

def PX(v):            # px -> Inches (96 px per inch)
    return Inches(v / 96.0)
def PTS(px):          # svg font px -> points
    return Pt(round(px * 0.75, 1))
def C(hexstr):
    return RGBColor.from_string(hexstr)

prs = Presentation()
prs.slide_width  = PX(1280)
prs.slide_height = PX(720)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

# dark background
bg = slide.background
bg.fill.solid()
bg.fill.fore_color.rgb = C("0F0F1A")

shapes = slide.shapes

def _set_dash(line):
    ln = line._get_or_add_ln()
    for e in ln.findall(qn('a:prstDash')):
        ln.remove(e)
    d = ln.makeelement(qn('a:prstDash'), {'val': 'dash'})
    ln.append(d)

def add_box(x, y, w, h, fill, line, lines, anchor=MSO_ANCHOR.TOP,
            rounded=True, dashed=False, align=PP_ALIGN.CENTER, lw=1.4):
    shp = shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        PX(x), PX(y), PX(w), PX(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    shp.line.color.rgb = C(line); shp.line.width = Pt(lw)
    if dashed:
        _set_dash(shp.line)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right'):
        setattr(tf, m, Inches(0.04))
    tf.margin_top = Inches(0.03); tf.margin_bottom = Inches(0.03)
    for i, (txt, sz, col, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = txt
        r.font.size = PTS(sz); r.font.bold = bold; r.font.color.rgb = C(col)
        r.font.name = 'Consolas'
    return shp

def add_label(cx, by, text, sz, col, w=150, align=PP_ALIGN.CENTER):
    tb = shapes.add_textbox(PX(cx - w / 2), PX(by - 12), PX(w), PX(16))
    tf = tb.text_frame; tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = PTS(sz); r.font.color.rgb = C(col); r.font.name = 'Consolas'
    return tb

def _arrow(conn, color, dashed, head=True, lw=1.7):
    conn.line.color.rgb = C(color); conn.line.width = Pt(lw)
    ln = conn.line._get_or_add_ln()
    if dashed:
        d = ln.makeelement(qn('a:prstDash'), {'val': 'dash'}); ln.append(d)
    if head:
        te = ln.makeelement(qn('a:tailEnd'),
                            {'type': 'triangle', 'w': 'med', 'len': 'med'})
        ln.append(te)

def add_arrow(x1, y1, x2, y2, color, dashed=False, lw=1.7):
    conn = shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                PX(x1), PX(y1), PX(x2), PX(y2))
    _arrow(conn, color, dashed, head=True, lw=lw)
    return conn

def add_route(points, color, dashed=False, lw=1.7):
    """multi-segment orthogonal route; arrowhead only on final segment."""
    for i in range(len(points) - 1):
        x1, y1 = points[i]; x2, y2 = points[i + 1]
        conn = shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    PX(x1), PX(y1), PX(x2), PX(y2))
        _arrow(conn, color, dashed, head=(i == len(points) - 2), lw=lw)

# ---------------- Title ----------------
tb = shapes.add_textbox(PX(36), PX(20), PX(820), PX(60))
tf = tb.text_frame; tf.word_wrap = False
p = tf.paragraphs[0]
r = p.add_run(); r.text = "CA — 模組訊號流 Module Signal Flow"
r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = C("E2E8F0"); r.font.name = 'Consolas'
p2 = tf.add_paragraph()
r2 = p2.add_run(); r2.text = "Conformer Accelerator · op ∈ {FFN, Conv, SHA, MHA}"
r2.font.size = Pt(10); r2.font.color.rgb = C("94A3B8"); r2.font.name = 'Consolas'

# ---------------- Legend ----------------
leg = add_box(946, 28, 300, 78, "0F172A", "334155",
              [("", 6, "94A3B8", False)], anchor=MSO_ANCHOR.TOP, lw=1.0)
ltf = leg.text_frame; ltf.clear()
legend_items = [("● data bus", "3B82F6"),
                ("● control / cmd", "F97316"),
                ("● result feedback", "EAB308"),
                ("● handshake", "64748B"),
                ("● write-back", "10B981"),
                ("● norm_data", "A855F7")]
for i, (txt, col) in enumerate(legend_items):
    p = ltf.paragraphs[0] if i == 0 else ltf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    r1 = p.add_run(); r1.text = txt[0]; r1.font.color.rgb = C(col); r1.font.size = Pt(8); r1.font.name = 'Consolas'
    r2 = p.add_run(); r2.text = txt[1:]; r2.font.color.rgb = C("94A3B8"); r2.font.size = Pt(8); r2.font.name = 'Consolas'

# ---------------- Module boxes ----------------
add_box(40, 150, 160, 420, "1E3A5F", "3B82F6",
        [("PATTERN", 16, "93C5FD", True),
         ("Testbench / host", 11, "94A3B8", False)], anchor=MSO_ANCHOR.TOP, lw=1.5)

add_box(250, 110, 270, 250, "1C1917", "F97316",
        [("CA_Control", 14, "FDBA74", True),
         ("FSM + RAM 排程", 10.5, "94A3B8", False)],
        anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, lw=1.5)

# FSM pills
pill = dict(fill="44403C", line="78716C")
for (px_, py_, pw_, txt) in [(262,176,52,"IDLE"), (330,176,96,"FAST_RUN"),
                             (262,210,98,"ATT_PARAM"), (374,210,52,"READ"),
                             (440,210,52,"WAIT")]:
    add_box(px_, py_, pw_, 22, pill["fill"], pill["line"],
            [(txt, 10, "FED7AA", False)], anchor=MSO_ANCHOR.MIDDLE, lw=1.0)
add_label(390, 191, "FFN/Conv", 9, "64748B", w=70, align=PP_ALIGN.LEFT)
add_label(390, 262, "▸ issue + 完成追蹤 (FAST/ATT 共用 writeback)",
          9.5, "64748B", w=260, align=PP_ALIGN.LEFT)

# RAM cylinder
ram = shapes.add_shape(MSO_SHAPE.CAN, PX(300), PX(410), PX(170), PX(155))
ram.fill.solid(); ram.fill.fore_color.rgb = C("052E16")
ram.line.color.rgb = C("10B981"); ram.line.width = Pt(1.5)
ram.shadow.inherit = False
rtf = ram.text_frame; rtf.word_wrap = True; rtf.vertical_anchor = MSO_ANCHOR.MIDDLE
for i,(txt,sz,col,bold) in enumerate([("RAM",16,"6EE7B7",True),
        ("256 × (8×8, 4b)",10.5,"94A3B8",False),
        ("burst 128 · neg-edge",9.5,"64748B",False)]):
    p = rtf.paragraphs[0] if i==0 else rtf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text=txt; r.font.size=PTS(sz); r.font.bold=bold
    r.font.color.rgb=C(col); r.font.name='Consolas'

# CA_DataPath container
add_box(580, 110, 660, 500, "14101F", "A855F7",
        [("CA_DataPath", 14, "C4B5FD", True),
         ("compute → activation → PoT", 10.5, "94A3B8", False)],
        anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, lw=1.5)

# ATT_Stream_Core group (dashed)
add_box(590, 250, 400, 300, "1A1530", "7C3AED",
        [("ATT_Stream_Core", 12, "C4B5FD", True)],
        anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, dashed=True, lw=1.2)

add_box(600, 300, 120, 180, "1E1B4B", "6366F1",
        [("QKV", 12, "C7D2FE", True), ("Proj + Quant", 9.5, "A5B4FC", False),
         ("Q = X·Wq", 9.5, "94A3B8", False), ("K = X·Wk", 9.5, "94A3B8", False),
         ("V = X·Wv", 9.5, "94A3B8", False), ("→ 4-bit", 9, "64748B", False)],
        anchor=MSO_ANCHOR.TOP)
add_box(734, 300, 120, 180, "1E1B4B", "6366F1",
        [("Score 8-Tap", 12, "C7D2FE", True), ("Q · Kᵀ", 12, "A5B4FC", False),
         ("SHA / MHA", 9.5, "94A3B8", False), ("col 切半 ×2", 9, "64748B", False)],
        anchor=MSO_ANCHOR.TOP)
add_box(868, 300, 116, 180, "1E1B4B", "6366F1",
        [("Final Acc", 12, "C7D2FE", True), ("act(s) · V", 12, "A5B4FC", False),
         ("Booth", 9.5, "94A3B8", False)],
        anchor=MSO_ANCHOR.TOP)
add_box(1000, 300, 100, 180, "1C1917", "F97316",
        [("ACT", 12, "FDBA74", True), ("ReLU / RAT", 9.5, "94A3B8", False),
         ("CAT / BAT", 9.5, "94A3B8", False), ("5-stage", 9, "64748B", False)],
        anchor=MSO_ANCHOR.TOP)
add_box(1114, 300, 110, 180, "1E1B4B", "A855F7",
        [("PoT Quant", 12, "D8B4FE", True), ("max → CLZ", 9.5, "94A3B8", False),
         ("shift · clamp", 9.5, "94A3B8", False), ("4-bit [-8,7]", 9, "64748B", False)],
        anchor=MSO_ANCHOR.TOP)

# ---------------- Arrows (signal flow) ----------------
add_arrow(200, 210, 250, 210, "F97316")
add_label(225, 202, "op·act·param", 9, "FB923C", w=90)
add_arrow(520, 180, 580, 180, "F97316")
add_label(550, 172, "issue·mode + exec_*", 9, "FB923C", w=130)
add_arrow(580, 232, 520, 232, "EAB308")
add_label(550, 225, "result_valid", 9, "EAB308", w=90)
add_arrow(358, 360, 358, 410, "F97316")
add_label(318, 392, "en·addr·burst", 9, "FB923C", w=80, align=PP_ALIGN.RIGHT)
add_arrow(412, 410, 412, 362, "64748B", dashed=True, lw=1.6)
add_label(458, 392, "ready·valid", 9, "94A3B8", w=80, align=PP_ALIGN.LEFT)
add_arrow(470, 455, 600, 455, "3B82F6", lw=2.0)
add_label(528, 448, "rd_data[255:0]", 9.5, "60A5FA", w=110)

add_route([(1169,480),(1169,650),(385,650),(385,587)], "10B981", dashed=True, lw=1.9)
add_label(770, 644, "wr_data[255:0]  (wr_valid)", 9.5, "34D399", w=260)
add_route([(1204,480),(1204,682),(120,682),(120,570)], "3B82F6", lw=1.9)
add_label(690, 676, "out_valid · out_data[31:0]  (last token)", 9.5, "60A5FA", w=330)

add_arrow(720, 390, 734, 390, "3B82F6", lw=1.8); add_label(727, 382, "q·k·v", 8.5, "60A5FA", w=44)
add_arrow(854, 390, 868, 390, "3B82F6", lw=1.8); add_label(861, 382, "score", 8.5, "60A5FA", w=44)
add_arrow(984, 390, 1000, 390, "3B82F6", lw=1.8); add_label(992, 382, "ctx", 8.5, "60A5FA", w=40)
add_arrow(1100, 390, 1114, 390, "3B82F6", lw=1.8); add_label(1107, 382, "act", 8.5, "60A5FA", w=40)

add_route([(660,480),(660,516),(1050,516),(1050,480)], "A855F7", dashed=True, lw=1.7)
add_label(852, 510, "norm_data  (FFN / Conv  y = W·x)", 9.5, "C084FC", w=240)

out = r"C:\Users\terry\OneDrive\\u6587件\\NYCU_DCS_2026\\DCS_FINAL\\report_figs\\ca_signalflow.pptx"
import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ca_signalflow.pptx")
prs.save(out)
print("Saved:", out)
print("shapes on slide:", len(slide.shapes._spTree))
